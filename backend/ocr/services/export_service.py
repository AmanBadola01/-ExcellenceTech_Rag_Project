"""
Export Service.
Exports extracted document data and line item tables to JSON, CSV, and formatted Excel (.xlsx).
"""

import io
import csv
import json
import zipfile
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


class ExportService:
    @staticmethod
    def to_json(doc_data: Dict[str, Any]) -> str:
        """Converts extracted document data to pretty JSON."""
        return json.dumps(doc_data, indent=2, ensure_ascii=False)

    @staticmethod
    def to_csv(fields: List[Dict[str, Any]], table: Dict[str, Any] = None) -> str:
        """Generates standard CSV output of fields and table rows."""
        output = io.StringIO()
        writer = csv.writer(output)

        # 1. Key-Value Fields
        writer.writerow(["Field", "Value", "Confidence"])
        for f in fields:
            writer.writerow([f.get("label", f.get("name")), f.get("value", ""), f"{f.get('confidence_pct', 0)}%"])

        # 2. Table Rows (if present)
        if table and table.get("rows"):
            writer.writerow([])
            writer.writerow([f"--- {table.get('name', 'Table')} ---"])
            headers = table.get("columns") or list(table["rows"][0].keys())
            writer.writerow(headers)
            for row in table["rows"]:
                writer.writerow([row.get(h, "") for h in headers])

        return output.getvalue()

    @staticmethod
    def to_excel(
        doc_name: str,
        doc_type: str,
        fields: List[Dict[str, Any]],
        table: Dict[str, Any] = None,
    ) -> bytes:
        """
        Generates professionally formatted .xlsx workbook with OpenPyXL.
        Sheet 1: Document Overview & Fields
        Sheet 2: Line Items / Tabular Data
        """
        wb = openpyxl.Workbook()
        ws_fields = wb.active
        ws_fields.title = "Extracted Fields"

        # Styling definitions
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Segoe UI", size=10)
        bold_font = Font(name="Segoe UI", size=10, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='E5E7EB'),
            right=Side(style='thin', color='E5E7EB'),
            top=Side(style='thin', color='E5E7EB'),
            bottom=Side(style='thin', color='E5E7EB'),
        )

        # Title Block
        ws_fields["A1"] = f"Extraction Report: {doc_name}"
        ws_fields["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
        ws_fields["A2"] = f"Document Type: {doc_type}"
        ws_fields["A2"].font = regular_font

        # Fields Table Headers
        headers = ["Field", "Extracted Value", "Category", "Confidence", "Discrepancy / Notes"]
        for col_idx, h_text in enumerate(headers, start=1):
            cell = ws_fields.cell(row=4, column=col_idx, value=h_text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="left", vertical="center")

        # Field Rows
        for r_idx, f in enumerate(fields, start=5):
            ws_fields.cell(row=r_idx, column=1, value=f.get("label", f.get("name"))).font = bold_font
            ws_fields.cell(row=r_idx, column=2, value=str(f.get("value", ""))).font = regular_font
            ws_fields.cell(row=r_idx, column=3, value=f.get("category", "")).font = regular_font
            ws_fields.cell(row=r_idx, column=4, value=f"{f.get('confidence_pct', 0)}%").font = regular_font
            ws_fields.cell(row=r_idx, column=5, value=f.get("discrepancy", "") or "OK").font = regular_font

            for c in range(1, 6):
                ws_fields.cell(row=r_idx, column=c).border = thin_border

        # Auto-adjust column widths
        for col in ws_fields.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws_fields.column_dimensions[col_letter].width = max(max_len + 4, 15)

        # Sheet 2: Tabular line items
        if table and table.get("rows"):
            ws_table = wb.create_sheet(title=table.get("name", "Table"))
            table_headers = table.get("columns") or list(table["rows"][0].keys())

            # Headers
            for col_idx, th in enumerate(table_headers, start=1):
                cell = ws_table.cell(row=1, column=col_idx, value=th)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center" if "qty" in th.lower() else "left")

            # Rows
            for row_idx, r_data in enumerate(table["rows"], start=2):
                for col_idx, th in enumerate(table_headers, start=1):
                    val = r_data.get(th, "")
                    c = ws_table.cell(row=row_idx, column=col_idx, value=val)
                    c.font = regular_font
                    c.border = thin_border

            for col in ws_table.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws_table.column_dimensions[col_letter].width = max(max_len + 4, 15)

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

